"""
Export per-site results to a spreadsheet: raw data plus a bar chart.

Produces both a plain CSV (for re-import anywhere) and an .xlsx with an
embedded chart - Google Drive preserves openpyxl charts when it converts an
uploaded .xlsx into a Google Sheet, so this is the file to hand to
create_file if you want the chart to show up there too.

Usage:
    python -m analysis.export_sheet
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

DATA_DIR = Path(__file__).parent.parent / "data"

QUERY = """
SELECT
  p.sector AS sector,
  p.url AS website,
  p.page_domain AS page_domain,
  p.status AS http_status,
  p.crawled_at AS crawled_at,
  count(r.request_id) AS total_requests,
  count(*) FILTER (WHERE r.third_party) AS third_party_requests,
  count(*) FILTER (WHERE r.is_tracker) AS tracker_requests,
  count(DISTINCT r.owner) FILTER (WHERE r.owner IS NOT NULL) AS distinct_tracker_owners,
  string_agg(DISTINCT r.owner, ', ') FILTER (WHERE r.owner IS NOT NULL) AS tracker_owners,
  count(*) FILTER (WHERE r.offshore) AS offshore_requests,
  string_agg(DISTINCT r.country, ', ') FILTER (WHERE r.offshore) AS offshore_countries,
  count(*) FILTER (WHERE r.insecure) AS insecure_requests,
  (SELECT count(*) FROM pii_findings f WHERE f.page_id = p.page_id) AS pii_findings
FROM pages p
LEFT JOIN requests r ON r.page_id = p.page_id
GROUP BY p.page_id, p.sector, p.url, p.page_domain, p.status, p.crawled_at
ORDER BY p.sector, p.url
"""


def fetch_rows(db_path: Path):
    con = duckdb.connect(str(db_path), read_only=True)
    cursor = con.execute(QUERY)
    columns = [d[0] for d in cursor.description]
    rows = cursor.fetchall()
    con.close()
    return columns, rows


def write_csv(columns: list[str], rows: list[tuple], out_path: Path) -> None:
    with out_path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        writer.writerows(rows)


def write_workbook(columns: list[str], rows: list[tuple], out_path: Path) -> None:
    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "per_site"
    data_ws.append(columns)
    for row in rows:
        data_ws.append(list(row))

    for i, column in enumerate(columns, start=1):
        widest = max([len(column)] + [len(str(r[i - 1])) for r in rows], default=len(column))
        data_ws.column_dimensions[get_column_letter(i)].width = min(widest + 2, 60)

    website_col = columns.index("website") + 1
    third_party_col = columns.index("third_party_requests") + 1
    tracker_col = columns.index("tracker_requests") + 1
    assert tracker_col == third_party_col + 1, "chart assumes these columns stay adjacent"
    n_rows = len(rows) + 1  # + header

    chart_ws = wb.create_sheet("chart")
    chart = BarChart()
    chart.type = "col"
    chart.title = "Third-party vs tracker requests per site"
    chart.y_axis.title = "Requests"
    chart.x_axis.title = "Website"
    chart.width = 32
    chart.height = 15

    data = Reference(
        data_ws, min_col=third_party_col, max_col=tracker_col, min_row=1, max_row=n_rows
    )
    categories = Reference(data_ws, min_col=website_col, min_row=2, max_row=n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart_ws.add_chart(chart, "B2")

    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export per-site results to CSV + chart workbook")
    parser.add_argument("--db", default=str(DATA_DIR / "dataexodus.duckdb"))
    parser.add_argument("--csv-out", default=str(DATA_DIR / "per_site.csv"))
    parser.add_argument("--xlsx-out", default=str(DATA_DIR / "per_site_report.xlsx"))
    args = parser.parse_args()

    db_path = Path(args.db)
    if not db_path.exists():
        raise SystemExit(f"{db_path} not found - run `python -m pipeline.process` first")

    columns, rows = fetch_rows(db_path)
    write_csv(columns, rows, Path(args.csv_out))
    write_workbook(columns, rows, Path(args.xlsx_out))
    print(f"{len(rows)} sites -> {args.csv_out}, {args.xlsx_out}")


if __name__ == "__main__":
    main()
